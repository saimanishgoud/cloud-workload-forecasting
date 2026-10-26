
from pathlib import Path

import matplotlib
matplotlib.use("Agg")

import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter


# Project paths
PROJECT_ROOT = Path(__file__).resolve().parents[2]
REPORT_DIR = PROJECT_ROOT / "reports"
CHART_DIR = REPORT_DIR / "charts"

TARGET_FILE = REPORT_DIR / "aggregate_targets.csv"
TRAIN_FILE = REPORT_DIR / "training_targets.csv"
HOLDOUT_FILE = REPORT_DIR / "holdout_timestamps.csv"

EXCEL_FILE = REPORT_DIR / "workload_analysis.xlsx"


def load_and_validate_targets():
    """Load and validate the aggregate target series."""

    if not TARGET_FILE.exists():
        raise FileNotFoundError(
            f"Target file not found: {TARGET_FILE}. "
            "Complete Step 4 first."
        )

    targets = pd.read_csv(TARGET_FILE)

    required = {
        "timestamp",
        "active_vm_count",
        "aggregate_cpu_utilization",
    }

    missing = required - set(targets.columns)

    if missing:
        raise ValueError(
            f"Missing target columns: {sorted(missing)}"
        )

    targets["timestamp"] = pd.to_numeric(
        targets["timestamp"], errors="coerce"
    )

    targets["active_vm_count"] = pd.to_numeric(
        targets["active_vm_count"], errors="coerce"
    )

    targets["aggregate_cpu_utilization"] = pd.to_numeric(
        targets["aggregate_cpu_utilization"], errors="coerce"
    )

    if targets.isna().any().any():
        raise ValueError(
            "Target file contains missing or invalid values."
        )

    if targets["timestamp"].duplicated().any():
        raise ValueError(
            "Duplicate timestamps found in target file."
        )

    targets = targets.sort_values("timestamp").reset_index(drop=True)

    return targets


def create_summary(targets):
    """Calculate descriptive statistics for both targets."""

    summary = pd.DataFrame({
        "metric": [
            "Number of timestamps",
            "Active VM count - mean",
            "Active VM count - median",
            "Active VM count - minimum",
            "Active VM count - maximum",
            "Active VM count - standard deviation",
            "Aggregate CPU - mean",
            "Aggregate CPU - median",
            "Aggregate CPU - minimum",
            "Aggregate CPU - maximum",
            "Aggregate CPU - standard deviation",
        ],
        "value": [
            len(targets),
            targets["active_vm_count"].mean(),
            targets["active_vm_count"].median(),
            targets["active_vm_count"].min(),
            targets["active_vm_count"].max(),
            targets["active_vm_count"].std(),
            targets["aggregate_cpu_utilization"].mean(),
            targets["aggregate_cpu_utilization"].median(),
            targets["aggregate_cpu_utilization"].min(),
            targets["aggregate_cpu_utilization"].max(),
            targets["aggregate_cpu_utilization"].std(),
        ],
    })

    return summary


def create_interval_report(targets):
    """Calculate the interval between consecutive timestamps."""

    intervals = targets[["timestamp"]].copy()

    intervals["interval_from_previous_seconds"] = (
        intervals["timestamp"].diff()
    )

    return intervals


def create_charts(targets):
    """Generate charts for workload activity and CPU utilization."""

    CHART_DIR.mkdir(parents=True, exist_ok=True)

    # Chart 1: Active VM count over time
    plt.figure(figsize=(12, 5))
    plt.plot(
        targets["timestamp"],
        targets["active_vm_count"],
    )
    plt.title("Active Virtual Machines Over Time")
    plt.xlabel("Timestamp (seconds relative to source trace)")
    plt.ylabel("Active VM Count")
    plt.grid(True, alpha=0.3)
    plt.tight_layout()
    plt.savefig(
        CHART_DIR / "active_vm_count.png",
        dpi=150,
    )
    plt.close()

    # Chart 2: Aggregate CPU utilization over time
    plt.figure(figsize=(12, 5))
    plt.plot(
        targets["timestamp"],
        targets["aggregate_cpu_utilization"],
    )
    plt.title("Aggregate CPU Utilization Over Time")
    plt.xlabel("Timestamp (seconds relative to source trace)")
    plt.ylabel("Mean CPU Utilization (%)")
    plt.grid(True, alpha=0.3)
    plt.tight_layout()
    plt.savefig(
        CHART_DIR / "aggregate_cpu_utilization.png",
        dpi=150,
    )
    plt.close()

    # Chart 3: Distribution of aggregate CPU utilization
    plt.figure(figsize=(9, 5))
    plt.hist(
        targets["aggregate_cpu_utilization"],
        bins=20,
        edgecolor="black",
    )
    plt.title("Distribution of Aggregate CPU Utilization")
    plt.xlabel("Mean CPU Utilization (%)")
    plt.ylabel("Number of Timestamps")
    plt.grid(True, axis="y", alpha=0.3)
    plt.tight_layout()
    plt.savefig(
        CHART_DIR / "cpu_utilization_distribution.png",
        dpi=150,
    )
    plt.close()

    # Chart 4: Change in active VM count
    changes = targets["active_vm_count"].diff()

    plt.figure(figsize=(12, 5))
    plt.plot(targets["timestamp"], changes)
    plt.axhline(0, linewidth=1)
    plt.title("Change in Active VM Count")
    plt.xlabel("Timestamp (seconds relative to source trace)")
    plt.ylabel("Change Since Previous Observation")
    plt.grid(True, alpha=0.3)
    plt.tight_layout()
    plt.savefig(
        CHART_DIR / "active_vm_count_change.png",
        dpi=150,
    )
    plt.close()


def format_workbook(file_path):
    """Format worksheet headers and add Excel line charts."""

    workbook = load_workbook(file_path)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )

        for column_cells in worksheet.iter_cols():
            column_letter = get_column_letter(
                column_cells[0].column
            )

            header = str(column_cells[0].value or "")

            worksheet.column_dimensions[
                column_letter
            ].width = min(max(len(header) + 3, 12), 32)

    # Add charts to the Targets worksheet.
    if "Targets" in workbook.sheetnames:
        ws = workbook["Targets"]

        if ws.max_row >= 3:
            chart = LineChart()
            chart.title = "Active VM Count"
            chart.y_axis.title = "Active VMs"
            chart.x_axis.title = "Timestamp"
            chart.height = 8
            chart.width = 16

            data = Reference(
                ws,
                min_col=2,
                min_row=1,
                max_row=ws.max_row,
            )

            categories = Reference(
                ws,
                min_col=1,
                min_row=2,
                max_row=ws.max_row,
            )

            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            ws.add_chart(chart, "E2")

            cpu_chart = LineChart()
            cpu_chart.title = "Aggregate CPU Utilization"
            cpu_chart.y_axis.title = "Mean CPU (%)"
            cpu_chart.x_axis.title = "Timestamp"
            cpu_chart.height = 8
            cpu_chart.width = 16

            cpu_data = Reference(
                ws,
                min_col=3,
                min_row=1,
                max_row=ws.max_row,
            )

            cpu_chart.add_data(
                cpu_data,
                titles_from_data=True,
            )
            cpu_chart.set_categories(categories)
            ws.add_chart(cpu_chart, "E18")

    workbook.save(file_path)


def main():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    CHART_DIR.mkdir(parents=True, exist_ok=True)

    targets = load_and_validate_targets()

    summary = create_summary(targets)
    intervals = create_interval_report(targets)

    # Export the target series, summary statistics,
    # and timestamp intervals to Excel.
    with pd.ExcelWriter(
        EXCEL_FILE,
        engine="openpyxl",
    ) as writer:
        targets.to_excel(
            writer,
            sheet_name="Targets",
            index=False,
        )

        summary.to_excel(
            writer,
            sheet_name="Summary",
            index=False,
        )

        intervals.to_excel(
            writer,
            sheet_name="Timestamp_Intervals",
            index=False,
        )

        # Include the existing training and holdout files.
        if TRAIN_FILE.exists():
            pd.read_csv(TRAIN_FILE).to_excel(
                writer,
                sheet_name="Training",
                index=False,
            )

        if HOLDOUT_FILE.exists():
            pd.read_csv(HOLDOUT_FILE).to_excel(
                writer,
                sheet_name="Holdout_Timestamps",
                index=False,
            )

    create_charts(targets)
    format_workbook(EXCEL_FILE)

    print("\nEXPORT COMPLETED SUCCESSFULLY")
    print(f"Excel workbook: {EXCEL_FILE}")
    print(f"Charts folder: {CHART_DIR}")

    print("\nSummary statistics:")
    print(summary.to_string(index=False))

    print("\nTimestamp interval counts:")
    print(
        intervals[
            "interval_from_previous_seconds"
        ].value_counts(dropna=False).head(10).to_string()
    )


if __name__ == "__main__":
    main()

